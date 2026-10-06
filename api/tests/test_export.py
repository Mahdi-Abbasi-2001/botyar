import io
import os

os.environ["DATABASE_URL"] = "sqlite:///./test.db"

import pytest  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

from app.catalog import parse_delimited, parse_xlsx  # noqa: E402
from app.db import Base, SessionLocal, engine  # noqa: E402
from app.main import app  # noqa: E402
from app.models import Bot, BotVersion, Product, Record, User  # noqa: E402

SPEC = {"name": "shop", "welcome": "سلام", "menu": [{"label": "خرید", "block": "shop"}],
        "blocks": [{"type": "catalog_order", "id": "shop", "title": "فروشگاه", "source": "table", "items": []}]}


@pytest.fixture()
def ctx():
    Base.metadata.drop_all(engine)
    Base.metadata.create_all(engine)
    with TestClient(app) as c:
        tok = c.post("/api/auth/register", json={"username": "ex_x.com", "password": "123456"}).json()["token"]
        H = {"Authorization": f"Bearer {tok}"}
        with SessionLocal() as db:
            uid = db.query(User).filter(User.username == "ex_x.com").one().id
            bot = Bot(user_id=uid, name="shop")
            db.add(bot)
            db.flush()
            db.add(BotVersion(bot_id=bot.id, version=1, spec=SPEC, note=""))
            db.add(Product(bot_id=bot.id, block_id="shop", name="کفش کتانی", category="ورزشی", price=2800000, stock=12,
                           options=[{"name": "سایز", "choices": ["۴۰", "۴۱"]}, {"name": "رنگ", "choices": ["مشکی", "سفید"]}], position=0))
            db.add(Product(bot_id=bot.id, block_id="shop", name="=HYPERLINK(\"http://evil\")", category="", price=1000, stock=None, options=[], position=1))
            db.add(Record(bot_id=bot.id, collection="shop", sandbox=False, data={
                "name": "مریم", "phone": "09123456789", "status": "new", "total": 980000,
                "items": [{"id": 1, "name": "کفش کتانی", "price": 490000, "options": {"سایز": "۴۰", "رنگ": "مشکی"}, "qty": 2}]}))
            db.add(Record(bot_id=bot.id, collection="shop", sandbox=False, data={"name": "+SUM(1+1)", "phone": "09120000000", "status": "new", "total": 5, "items": []}))
            db.commit()
            bid = bot.id
        yield c, H, bid


def test_products_csv_has_bom_and_round_trips_through_our_importer(ctx):
    c, H, bid = ctx
    r = c.get(f"/api/bots/{bid}/export/products?format=csv", headers=H)
    assert r.status_code == 200 and r.content.startswith("﻿".encode()) and "attachment" in r.headers["content-disposition"]
    rows = parse_delimited(r.content.decode("utf-8-sig"))
    assert rows[0] == ["نام", "دسته", "قیمت (تومان)", "موجودی", "سایز", "رنگ", "توضیحات"]
    assert rows[1][:5] == ["کفش کتانی", "ورزشی", "2800000", "12", "۴۰ / ۴۱"]


def test_products_xlsx_is_readable_by_the_importer(ctx):
    c, H, bid = ctx
    r = c.get(f"/api/bots/{bid}/export/products?format=xlsx", headers=H)
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
    rows = parse_xlsx(r.content)
    assert rows[1][0] == "کفش کتانی" and rows[1][2] == "2800000"


def test_formula_injection_is_neutralised_in_every_format(ctx):
    c, H, bid = ctx
    csv_prod = c.get(f"/api/bots/{bid}/export/products?format=csv", headers=H).content.decode("utf-8-sig")
    assert "'=HYPERLINK" in csv_prod and ',=HYPERLINK' not in csv_prod
    csv_rec = c.get(f"/api/bots/{bid}/export/records?format=csv", headers=H).content.decode("utf-8-sig")
    assert "'+SUM(1+1)" in csv_rec
    ws = load_workbook(io.BytesIO(c.get(f"/api/bots/{bid}/export/records?format=xlsx", headers=H).content)).active
    assert any(str(cell.value).startswith("'+SUM") for row in ws.iter_rows() for cell in row)


def test_records_xlsx_keeps_phone_numbers_as_text_and_describes_orders(ctx):
    c, H, bid = ctx
    ws = load_workbook(io.BytesIO(c.get(f"/api/bots/{bid}/export/records?format=xlsx", headers=H).content)).active
    header = [x.value for x in ws[1]]
    assert header[:3] == ["شناسه", "تاریخ", "بخش"] and "موبایل" in header and "اقلام" in header and "جمع (تومان)" in header
    row = {h: ws.cell(row=2, column=i + 1) for i, h in enumerate(header)}
    assert row["موبایل"].value == "09123456789" and row["موبایل"].data_type == "s"  # leading zero survives
    assert row["اقلام"].value == "کفش کتانی ۴۰ مشکی × 2" and row["وضعیت"].value == "جدید" and row["بخش"].value == "فروشگاه"
    assert ws.sheet_view.rightToLeft is True


def test_access_control_and_errors(ctx):
    c, H, bid = ctx
    other = c.post("/api/auth/register", json={"username": "other_x.com", "password": "123456"}).json()["token"]
    O = {"Authorization": f"Bearer {other}"}
    assert c.get(f"/api/bots/{bid}/export/products", headers=O).status_code == 404
    assert c.get(f"/api/bots/{bid}/export/records", headers=O).status_code == 404
    assert c.get(f"/api/bots/{bid}/export/products").status_code == 401
    assert c.get(f"/api/bots/{bid}/export/products?format=pdf", headers=H).status_code == 400
    assert c.get(f"/api/bots/{bid}/export/records?sandbox=true", headers=H).status_code == 404  # no sandbox records -> clear error
