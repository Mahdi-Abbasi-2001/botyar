"""FAQ through the real HTTP/Bale layers (embeddings and the variant-writing LLM are faked; no network)."""
import io
import os

os.environ["DATABASE_URL"] = "sqlite:///./test.db"

import numpy as np  # noqa: E402
import pytest  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

from app import bale, faq_index, faq_match  # noqa: E402
from app.config import settings  # noqa: E402
from app.db import Base, SessionLocal, engine  # noqa: E402
from app.main import app  # noqa: E402
from app.models import Bot, BotVersion, Record, User  # noqa: E402
from app.spec import BotSpec  # noqa: E402

ENTRIES = [{"question": "ساعت کاری کلینیک چیست؟", "answer": "شنبه تا چهارشنبه از ۹ صبح تا ۶ عصر."},
           {"question": "آدرس کلینیک کجاست؟", "answer": "خیابان ولیعصر، پلاک ۱۲."},
           {"question": "هزینه ویزیت چقدر است؟", "answer": "ویزیت اولیه ۲۵۰ هزار تومان است."}]
SPEC = {"name": "کلینیک", "welcome": "سلام", "menu": [{"label": "سؤال دارید؟", "block": "faq"}],
        "blocks": [{"type": "faq", "id": "faq", "title": "پرسش‌ها", "entries": ENTRIES},
                   {"type": "admin_notify", "id": "n", "on": "faq", "text": "سؤال جدید"}]}


def vec(texts, dim=128):
    out = np.zeros((len(texts), dim), dtype=np.float32)
    for r, t in enumerate(texts):
        for w in faq_match.tokens(t):
            out[r, hash(w) % dim] += 1.0
    n = np.linalg.norm(out, axis=1, keepdims=True)
    n[n == 0] = 1
    return out / n


@pytest.fixture()
def world(monkeypatch):
    Base.metadata.drop_all(engine)
    Base.metadata.create_all(engine)
    faq_index._cache.clear()
    faq_index._hits.clear()
    monkeypatch.setattr(settings, "public_base_url", "https://example.test")
    monkeypatch.setattr(settings, "bale_shared_bot_token", "SHAREDTOKEN")
    monkeypatch.setattr(faq_index, "_embed", lambda texts, timeout=6.0: (vec(texts), 10))
    monkeypatch.setattr(faq_index, "_variants", lambda db, bot_id, run_id, batch: {i: [f"میشه بگید {e.question}"] for i, e in batch})
    sent = []

    def fake(token, method, payload=None, timeout=15):
        if method == "sendMessage":
            sent.append((str(payload["chat_id"]), payload["text"], payload.get("reply_markup")))
        return {"username": "botyar_test_bot"} if method == "getMe" else True

    monkeypatch.setattr(bale, "api_call", fake)
    with TestClient(app) as c:
        tok = c.post("/api/auth/register", json={"username": "faq_x.com", "password": "123456"}).json()["token"]
        H = {"Authorization": f"Bearer {tok}"}
        with SessionLocal() as db:
            uid = db.query(User).filter(User.username == "faq_x.com").one().id
            bot = Bot(user_id=uid, name="کلینیک")
            db.add(bot)
            db.flush()
            db.add(BotVersion(bot_id=bot.id, version=1, spec=SPEC, note=""))
            db.commit()
            bid = bot.id
        yield c, H, bid, sent


def sim(c, H, bid, *texts, sid="s"):
    out = None
    for t in texts:
        out = c.post(f"/api/bots/{bid}/simulate", json={"session_id": sid, "text": t}, headers=H).json()["actions"]
    return out


def test_simulator_answers_through_the_embedding_index_and_publishing_builds_it(world):
    c, H, bid, sent = world
    from app.models import FaqIndex

    assert c.post(f"/api/bots/{bid}/publish", json={"mode": "shared"}, headers=H).status_code == 200   # publish indexes the FAQ
    with SessionLocal() as db:
        assert db.query(FaqIndex).filter(FaqIndex.bot_id == bid).count() == 6                         # 3 questions + 3 variants
    out = sim(c, H, bid, "/start", "m:0", "میشه بگید آدرس کلینیک کجاست؟")
    assert out[0]["text"] == "خیابان ولیعصر، پلاک ۱۲."


def test_without_an_index_the_bot_still_answers_by_word_matching(world):
    c, H, bid, sent = world
    out = sim(c, H, bid, "/start", "m:0", "آدرس کلینیک کجاست؟")                                          # nothing indexed yet
    assert out[0]["text"] == "خیابان ولیعصر، پلاک ۱۲."


def test_bale_unanswered_question_notifies_the_owner_and_the_owner_can_mark_it_handled(world):
    c, H, bid, sent = world
    st = c.post(f"/api/bots/{bid}/publish", json={"mode": "shared"}, headers=H).json()
    url = f"/api/hook/shared/{bale.shared_hook_secret()}"

    counter = [5000]

    def msg(chat, text):
        counter[0] += 1
        return {"update_id": counter[0], "message": {"message_id": 1, "chat": {"id": chat, "type": "private"}, "text": text}}

    def tap(chat, data):
        counter[0] += 1
        return {"update_id": counter[0], "callback_query": {"id": "q", "from": {"id": chat}, "data": data, "message": {"message_id": 7, "chat": {"id": chat, "type": "private"}}}}

    c.post(url, json=msg(900, f"/admin {st['admin_code']}"))
    c.post(url, json=msg(901, f"/start {st['code']}"))
    c.post(url, json=tap(901, "m:0"))
    sent.clear()
    c.post(url, json=msg(901, "پیتزا سفارش میدم"))
    owner_msgs = [t for chat, t, _ in sent if chat == "900"]
    assert any("❓ سؤال بدون پاسخ" in t and "پیتزا" in t for t in owner_msgs)
    assert any("پیدا نکردم" in t for chat, t, _ in sent if chat == "901")

    recs = c.get(f"/api/bots/{bid}/records?sandbox=false", headers=H).json()
    assert len(recs) == 1 and recs[0]["data"]["status"] == "unanswered" and not any(k.startswith("_") for k in recs[0]["data"])
    rid = recs[0]["id"]
    assert c.patch(f"/api/bots/{bid}/records/{rid}", json={"action": "status", "status": "handled"}, headers=H).json()["status"] == "handled"
    assert c.patch(f"/api/bots/{bid}/records/{rid}", json={"action": "status", "status": "handled"}, headers=H).status_code == 409
    assert c.patch(f"/api/bots/{bid}/records/{rid}", json={"action": "cancel"}, headers=H).status_code == 409
    assert c.patch(f"/api/bots/{bid}/records/{rid}", json={"action": "status", "status": "ready"}, headers=H).status_code == 409


def test_unanswered_questions_export_with_persian_headers(world):
    c, H, bid, sent = world
    sim(c, H, bid, "/start", "m:0", "قیمت دلار امروز چنده")
    ws = load_workbook(io.BytesIO(c.get(f"/api/bots/{bid}/export/records?format=xlsx&sandbox=true", headers=H).content)).active
    header = [x.value for x in ws[1]]
    row = {h: ws.cell(row=2, column=i + 1).value for i, h in enumerate(header)}
    assert "سؤال" in header and row["سؤال"] == "قیمت دلار امروز چنده" and row["وضعیت"] == "بدون پاسخ"


def test_faq_spec_limits():
    many = [{"question": f"سؤال شماره {i}", "answer": "پاسخ"} for i in range(41)]
    with pytest.raises(ValueError):
        BotSpec.model_validate({**SPEC, "blocks": [{"type": "faq", "id": "faq", "title": "t", "entries": many}]})
    with pytest.raises(ValueError):
        BotSpec.model_validate({**SPEC, "blocks": [{"type": "faq", "id": "faq", "title": "t", "entries": []}]})
