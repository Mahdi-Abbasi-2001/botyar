"""Catalog (product table) API: CRUD + import from CSV / Excel / pasted table (the agent maps the columns,
plain code converts every row) and from photos / PDF pages (vision extraction)."""
from __future__ import annotations

import base64
import csv
import io
import logging
import re
from datetime import datetime, timedelta, timezone
from typing import Literal

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from pydantic import BaseModel, Field, model_validator
from sqlalchemy import delete, func, select
from sqlalchemy.orm import Session

from . import engine, llm
from .auth import current_user
from .config import settings
from .db import get_db
from .models import Bot, BotFile, BotVersion, LlmCall, Product, User
from .spec import ProductIn, ProductOption

log = logging.getLogger("botyar.catalog")
router = APIRouter()
MAX_ROWS, MAX_FILE, MAX_IMAGES, DAILY_IMPORTS = 3000, 8_000_000, 4, 30


# ---------- helpers ----------
def _own(bot_id: int, user: User, db: Session) -> Bot:
    bot = db.get(Bot, bot_id)
    if not bot or bot.user_id != user.id:
        raise HTTPException(404, "ربات یافت نشد")
    return bot


def table_blocks(bot_id: int, db: Session) -> list[dict]:
    v = db.scalars(select(BotVersion).where(BotVersion.bot_id == bot_id).order_by(BotVersion.version.desc())).first()
    if v is None:
        return []
    return [{"id": b["id"], "title": b.get("title", b["id"])} for b in v.spec["blocks"] if b["type"] == "catalog_order" and b.get("source") == "table"]


def _block(bot_id: int, block: str | None, db: Session) -> str:
    blocks = table_blocks(bot_id, db)
    if not blocks:
        raise HTTPException(409, "این ربات بخش فروشگاهی با فهرست محصولات ندارد")
    if block is None:
        return blocks[0]["id"]
    if block not in [b["id"] for b in blocks]:
        raise HTTPException(404, "بخش یافت نشد")
    return block


def _prod_out(p: Product) -> dict:
    return {"id": p.id, "name": p.name, "category": p.category, "price": p.price, "stock": p.stock, "options": p.options or [],
            "description": p.description, "is_sample": p.is_sample, "photo": bool(p.has_photo)}


# ---------- parsing ----------
def parse_number(s) -> int | None:
    """'۱,۲۵۰,۰۰۰' -> 1250000; '450 هزار تومان' -> 450000; '' -> None."""
    t = engine.norm(str(s if s is not None else "")).replace("٬", ",").replace("٫", ".").replace("،", ",")
    m = re.search(r"\d[\d,]*(?:\.\d+)?", t)
    if not m:
        return None
    raw = m.group(0)
    val = float(raw.replace(",", "")) if re.fullmatch(r"\d+\.\d{1,2}", raw.replace(",", "")) else float(re.sub(r"[,.]", "", raw))
    if "میلیون" in t:
        val *= 1_000_000
    elif "هزار" in t:
        val *= 1_000
    return int(round(val))


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-16", "cp1256"):
        try:
            return data.decode(enc)
        except UnicodeError:
            continue
    return data.decode("utf-8", errors="replace")


def parse_delimited(text: str) -> list[list[str]]:
    sample = text[:4000]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel_tab if "\t" in sample else csv.excel
    rows = [[c.strip() for c in r] for r in csv.reader(io.StringIO(text), dialect)]
    return [r for r in rows if any(r)]


def parse_xlsx(data: bytes) -> list[list[str]]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = max(wb.worksheets, key=lambda w: (w.max_row or 0) * (w.max_column or 0))
    out = []
    for row in ws.iter_rows(values_only=True):
        cells = [("" if v is None else (str(int(v)) if isinstance(v, float) and v.is_integer() else str(v))).strip() for v in row]
        if any(cells):
            out.append(cells)
        if len(out) > MAX_ROWS + 20:
            break
    return out


_P2A = str.maketrans("۰۱۲۳۴۵۶۷۸۹", "0123456789")
_RANGE = re.compile(r"^(\D*?)(\d+)\s*(?:تا|to|-|–|—)\s*(\d+)(\D*)$")


def expand_ranges(choices: list[str]) -> list[str]:
    """«۴۰ تا ۴۴» -> ۴۰، ۴۱، ۴۲، ۴۳، ۴۴ (keeps the digit script). Only short, sane ranges are expanded."""
    out: list[str] = []
    for c in choices:
        m = _RANGE.match(c.strip().translate(_P2A))
        if m and 0 < int(m.group(3)) - int(m.group(2)) <= 30:
            persian = any(ch in "۰۱۲۳۴۵۶۷۸۹" for ch in c)
            for n in range(int(m.group(2)), int(m.group(3)) + 1):
                t = str(n)
                out.append(m.group(1).strip() + (t.translate(str.maketrans("0123456789", "۰۱۲۳۴۵۶۷۸۹")) if persian else t) + m.group(4))
        else:
            out.append(c)
    return list(dict.fromkeys(x.strip() for x in out if x.strip()))


def clean_options(options: list[dict]) -> list[dict]:
    res = []
    for o in options:
        ch = expand_ranges(o.get("choices", []))
        if ch:
            res.append({"name": o["name"], "choices": ch})
    return res


# ---------- agent: column mapping ----------
class OptionCol(BaseModel):
    group: str      # e.g. «سایز»
    column: int
    separator: str  # what separates several values inside one cell, e.g. "/" or "،"


class ColumnMapping(BaseModel):
    header_row: int | None       # row index of the header line, null when the table has none
    name_col: int
    category_col: int | None
    price_col: int
    price_unit: Literal["toman", "rial", "thousand_toman"]
    stock_col: int | None
    description_col: int | None
    option_cols: list[OptionCol]
    notes: list[str]             # short Persian remarks for the owner (assumptions, doubts)


MAP_PROMPT = """You map a store owner's product spreadsheet to Botyar's product table. Input: the first rows of a table, each row prefixed with its row index and every cell with its column index, like  r3: [0]=پیراهن آبی | [1]=مردانه | [2]=450000 .
Return: header_row (index of the header line or null), name_col, category_col, price_col, price_unit, stock_col, description_col and option_cols (columns that list choices like sizes or colors in one cell; give the group name in Persian and the separator used inside the cell).
Rules: use null for a column that does not exist; never invent columns. Iranian stores list prices in toman unless the header says ریال (rial) or «هزار تومان»/thousand (thousand_toman); if unsure choose toman. A column of sizes/colors is an option column, NOT the name. `notes`: at most 3 short Persian remarks for the owner about assumptions or doubts (empty list if none)."""


def map_columns(db: Session, bot_id: int, rows: list[list[str]]) -> ColumnMapping:
    sample = rows[:12]
    text = "\n".join(f"r{i}: " + " | ".join(f"[{j}]={c[:60]}" for j, c in enumerate(r)) for i, r in enumerate(sample))
    m: ColumnMapping = llm.call(db, bot_id=bot_id, run_id=0, step="catalog_map", instructions=MAP_PROMPT, input=text, schema=ColumnMapping, effort="low")
    width = max(len(r) for r in rows)
    for name in ("name_col", "price_col"):
        if not 0 <= getattr(m, name) < width:
            raise HTTPException(422, "ستون‌های فایل تشخیص داده نشدند؛ مطمئن شوید فایل ستون‌های «نام» و «قیمت» را دارد")
    return m


def split_choices(cell: str, sep: str) -> list[str]:
    """Split a cell of choices on literal separators. A plain word such as «تا» is a range marker, never a separator."""
    seps = ["/", "،", ",", "|", ";", "؛"]
    sep = (sep or "").strip()
    if sep and sep not in seps and not (len(sep) > 1 and sep.isalpha()):
        seps.append(sep)
    return [v.strip() for v in re.split("|".join(re.escape(x) for x in seps), cell) if v.strip()]


def _cell(r: list[str], i: int | None) -> str:
    return r[i].strip() if i is not None and 0 <= i < len(r) else ""


def convert(rows: list[list[str]], m: ColumnMapping) -> tuple[list[dict], list[str]]:
    start = (m.header_row + 1) if m.header_row is not None else 0
    out, no_price, warnings = [], 0, []
    mult = {"toman": 1, "rial": 0.1, "thousand_toman": 1000}[m.price_unit]
    for r in rows[start:]:
        name = _cell(r, m.name_col)
        if not name:
            continue
        price = parse_number(_cell(r, m.price_col))
        if price is None:
            no_price += 1
            continue
        stock_raw = _cell(r, m.stock_col)
        if not stock_raw:
            stock = None
        elif any(w in stock_raw for w in ("ناموجود", "تمام", "اتمام")):
            stock = 0
        else:
            stock = parse_number(stock_raw)
        options = []
        for oc in m.option_cols:
            vals = expand_ranges(split_choices(_cell(r, oc.column), oc.separator))
            if vals:
                options.append({"name": oc.group, "choices": vals})
        out.append({"name": name[:300], "category": _cell(r, m.category_col)[:120], "price": int(round(price * mult)),
                    "stock": stock, "options": options, "description": _cell(r, m.description_col)[:1000]})
        if len(out) >= MAX_ROWS:
            warnings.append(f"فقط {MAX_ROWS} محصول اول خوانده شد.")
            break
    if no_price:
        warnings.append(f"{no_price} ردیف بدون قیمت معتبر نادیده گرفته شد.")
    return out, warnings


# ---------- agent: vision ----------
class Extracted(BaseModel):
    products: list[ProductIn]
    warnings: list[str]


VISION_PROMPT = """You read photos or scanned pages of a store's price list / menu / catalog (usually Persian) and extract every product you can read.
Rules: copy names exactly as written; price as an integer in TOMAN (if the page clearly says ریال divide by 10; «هزار تومان» multiply by 1000); category from the section heading when there is one, otherwise empty string; stock null unless written; options only when sizes/colors are listed for that product; description empty unless written. Do NOT invent products or prices — skip anything unreadable and say so in `warnings` (short Persian sentences, at most 3)."""


def pdf_to_images(data: bytes) -> list[bytes]:
    import pymupdf

    try:
        doc = pymupdf.open(stream=data, filetype="pdf")
        pages = [doc[i].get_pixmap(dpi=110).tobytes("png") for i in range(min(len(doc), MAX_IMAGES))]
    except Exception:  # noqa: BLE001 - corrupt / encrypted / not a PDF
        raise HTTPException(422, "فایل PDF خوانده نشد؛ فایل دیگری را امتحان کنید")
    if not pages:
        raise HTTPException(422, "فایل PDF صفحه‌ای ندارد")
    return pages


def checked_image(data: bytes) -> tuple[bytes, str]:
    """Reject files that are not real images before spending a model call on them."""
    from PIL import Image

    try:
        with Image.open(io.BytesIO(data)) as im:
            fmt = (im.format or "").lower()
            im.verify()
    except Exception:  # noqa: BLE001
        raise HTTPException(422, "تصویر خوانده نشد؛ عکس یا تصویر صفحه‌ی واضح‌تری بارگذاری کنید")
    mime = {"png": "image/png", "jpeg": "image/jpeg", "webp": "image/webp", "gif": "image/gif"}.get(fmt)
    if mime is None:
        raise HTTPException(422, "این قالب تصویر پشتیبانی نمی‌شود؛ تصویر PNG یا JPG بارگذاری کنید")
    return data, mime


def extract_from_images(db: Session, bot_id: int, images: list[tuple[bytes, str]]) -> Extracted:
    parts: list[dict] = [{"type": "input_text", "text": "Extract the products from these pages."}]
    for data, mime in images:
        parts.append({"type": "input_image", "image_url": f"data:{mime};base64,{base64.b64encode(data).decode()}"})
    return llm.call(db, bot_id=bot_id, run_id=0, step="catalog_vision", instructions=VISION_PROMPT,
                    input=[{"role": "user", "content": parts}], schema=Extracted, effort="low", max_out=16000)


# ---------- endpoints ----------
def _imports_today(user: User, db: Session) -> int:
    since = datetime.now(timezone.utc) - timedelta(hours=24)
    return db.scalar(select(func.count()).select_from(LlmCall).join(Bot, Bot.id == LlmCall.bot_id).where(
        Bot.user_id == user.id, LlmCall.step.in_(["catalog_map", "catalog_vision"]), LlmCall.created_at >= since)) or 0


@router.get("/api/bots/{bot_id}/catalog")
def get_catalog(bot_id: int, block: str | None = None, user: User = Depends(current_user), db: Session = Depends(get_db)):
    _own(bot_id, user, db)
    blocks = table_blocks(bot_id, db)
    if not blocks:
        return {"blocks": [], "block": None, "total": 0, "products": [], "categories": [], "sample": False}
    b = _block(bot_id, block, db)
    rows = list(db.scalars(select(Product).where(Product.bot_id == bot_id, Product.block_id == b).order_by(Product.position, Product.id)))
    return {"blocks": blocks, "block": b, "total": len(rows), "products": [_prod_out(p) for p in rows[:300]],
            "categories": engine.catalog_categories([{"category": p.category} for p in rows]),
            "sample": bool(rows) and all(p.is_sample for p in rows)}


@router.post("/api/bots/{bot_id}/catalog/preview")
async def preview(bot_id: int, files: list[UploadFile] = File(default=[]), text: str = Form(""), user: User = Depends(current_user),
                  db: Session = Depends(get_db)):
    _own(bot_id, user, db)
    _block(bot_id, None, db)
    if _imports_today(user, db) >= DAILY_IMPORTS:
        raise HTTPException(429, "سقف واردسازی روزانه پر شده است")
    since = datetime.now(timezone.utc) - timedelta(hours=24)
    if (db.scalar(select(func.count()).select_from(LlmCall).where(LlmCall.step.in_(["catalog_map", "catalog_vision"]), LlmCall.created_at >= since)) or 0) >= settings.global_daily_imports:
        raise HTTPException(503, "ظرفیت امروز واردسازی تکمیل شده است؛ فردا دوباره تلاش کنید")
    if not files and not text.strip():
        raise HTTPException(400, "یک فایل انتخاب کنید یا جدول را جای‌گذاری کنید")
    from_vision = bool(files) and any((f.content_type or "").startswith("image/") or (f.filename or "").lower().endswith((".pdf", ".png", ".jpg", ".jpeg", ".webp")) for f in files)
    before = float(db.scalar(select(func.coalesce(func.sum(LlmCall.cost_usd), 0.0)).where(LlmCall.bot_id == bot_id)))
    notes: list[str] = []

    if from_vision:
        images: list[tuple[bytes, str]] = []
        for f in files:
            data = await f.read()
            if len(data) > MAX_FILE:
                raise HTTPException(413, "حجم فایل بیش از حد مجاز است")
            name = (f.filename or "").lower()
            if name.endswith(".pdf") or f.content_type == "application/pdf":
                images += [(im, "image/png") for im in pdf_to_images(data)]
            else:
                images.append(checked_image(data))
        images = images[:MAX_IMAGES]
        try:
            ex = extract_from_images(db, bot_id, images)
        except HTTPException:
            raise
        except Exception:  # noqa: BLE001 - provider error, timeout, unreadable image...
            log.exception("vision import failed")
            raise HTTPException(502, "خواندن تصویر ناموفق بود؛ چند لحظه بعد دوباره تلاش کنید")
        products, warnings, kind = [p.model_dump() for p in ex.products][:MAX_ROWS], list(ex.warnings), "vision"
        for pr in products:
            pr["options"] = clean_options(pr["options"])
        if not products:
            warnings.append("محصولی در تصویر پیدا نشد.")
    else:
        if files:
            f = files[0]
            data = await f.read()
            if len(data) > MAX_FILE:
                raise HTTPException(413, "حجم فایل بیش از حد مجاز است")
            name = (f.filename or "").lower()
            try:
                rows = parse_xlsx(data) if name.endswith((".xlsx", ".xlsm")) else parse_delimited(_decode(data))
            except Exception:  # noqa: BLE001
                raise HTTPException(422, "فایل خوانده نشد؛ فرمت CSV یا Excel (xlsx) را امتحان کنید")
        else:
            rows = parse_delimited(text)
        if len(rows) < 2:
            raise HTTPException(422, "جدول باید حداقل یک سطر عنوان و یک سطر محصول داشته باشد")
        try:
            m = map_columns(db, bot_id, rows)
        except HTTPException:
            raise
        except Exception:  # noqa: BLE001
            log.exception("column mapping failed")
            raise HTTPException(502, "تشخیص ستون‌ها ناموفق بود؛ چند لحظه بعد دوباره تلاش کنید")
        products, warnings = convert(rows, m)
        notes, kind = m.notes, "table"
    after = float(db.scalar(select(func.coalesce(func.sum(LlmCall.cost_usd), 0.0)).where(LlmCall.bot_id == bot_id)))
    return {"kind": kind, "total": len(products), "products": products, "warnings": warnings, "notes": notes, "cost_usd": after - before}


class CommitIn(BaseModel):
    block: str | None = None
    mode: Literal["replace", "append"] = "replace"
    products: list[ProductIn] = Field(max_length=MAX_ROWS)

    @model_validator(mode="after")
    def _sizes(self):
        for i, p in enumerate(self.products, 1):
            if not p.name.strip() or len(p.name) > 300 or len(p.category) > 120 or len(p.description) > 1000:
                raise ValueError(f"محصول {i}: نام باید ۱ تا ۳۰۰، دسته حداکثر ۱۲۰ و توضیح حداکثر ۱۰۰۰ نویسه باشد")
            if len(p.options) > 10 or any(len(o.name) > 60 or len(o.choices) > 60 or any(len(c) > 80 for c in o.choices) for o in p.options):
                raise ValueError(f"محصول {i}: گزینه‌ها بیش از حد بلند یا زیاد است")
        return self


@router.post("/api/bots/{bot_id}/catalog/commit")
def commit(bot_id: int, body: CommitIn, user: User = Depends(current_user), db: Session = Depends(get_db)):
    _own(bot_id, user, db)
    b = _block(bot_id, body.block, db)
    if not body.products:
        raise HTTPException(400, "فهرستی برای ذخیره وجود ندارد")
    existing = list(db.scalars(select(Product).where(Product.bot_id == bot_id, Product.block_id == b)))
    if body.mode == "replace" or all(p.is_sample for p in existing):  # demo rows never survive a real import
        _drop_photos(db, bot_id, [p.id for p in existing])
        db.execute(delete(Product).where(Product.bot_id == bot_id, Product.block_id == b))
        start = 0
    else:
        start = max((p.position for p in existing), default=-1) + 1
    for i, p in enumerate(body.products):
        db.add(Product(bot_id=bot_id, block_id=b, name=p.name, category=p.category, price=p.price, stock=p.stock,
                       options=[o.model_dump() for o in p.options if o.choices], description=p.description, position=start + i))
    db.commit()
    return {"ok": True, "saved": len(body.products)}


class ProductPatch(BaseModel):
    name: str | None = None
    category: str | None = None
    price: int | None = Field(default=None, ge=0)
    stock: int | None = Field(default=None, ge=0)
    clear_stock: bool = False  # unlimited


@router.patch("/api/bots/{bot_id}/catalog/products/{pid}")
def patch_product(bot_id: int, pid: int, body: ProductPatch, user: User = Depends(current_user), db: Session = Depends(get_db)):
    _own(bot_id, user, db)
    p = db.get(Product, pid)
    if not p or p.bot_id != bot_id:
        raise HTTPException(404, "محصول یافت نشد")
    if body.name is not None:
        p.name = body.name.strip()[:300]
    if body.category is not None:
        p.category = body.category.strip()[:120]
    if body.price is not None:
        p.price = body.price
    if body.clear_stock:
        p.stock = None
    elif body.stock is not None:
        p.stock = body.stock
    p.is_sample = False
    db.commit()
    return _prod_out(p)


@router.delete("/api/bots/{bot_id}/catalog/products/{pid}")
def delete_product(bot_id: int, pid: int, user: User = Depends(current_user), db: Session = Depends(get_db)):
    _own(bot_id, user, db)
    p = db.get(Product, pid)
    if not p or p.bot_id != bot_id:
        raise HTTPException(404, "محصول یافت نشد")
    _drop_photos(db, bot_id, [p.id])
    db.delete(p)
    db.commit()
    return {"ok": True}


# ---------- product photos (sent to the customer when they open the product) ----------
PHOTO_MAX = 1_500_000  # the panel shrinks photos before upload; this only stops oversized files


def _photo_key(pid: int) -> str:
    return f"product:{pid}"


def _drop_photos(db: Session, bot_id: int, pids: list[int]):
    if pids:
        db.execute(delete(BotFile).where(BotFile.bot_id == bot_id, BotFile.block_id.in_([_photo_key(i) for i in pids])))


def _product(bot_id: int, pid: int, db: Session) -> Product:
    p = db.get(Product, pid)
    if not p or p.bot_id != bot_id:
        raise HTTPException(404, "محصول یافت نشد")
    return p


@router.put("/api/bots/{bot_id}/catalog/products/{pid}/photo")
async def upload_photo(bot_id: int, pid: int, file: UploadFile = File(...), user: User = Depends(current_user), db: Session = Depends(get_db)):
    from .media import IMAGE_MAGIC, MAX_TOTAL
    _own(bot_id, user, db)
    p = _product(bot_id, pid, db)
    data = await file.read(PHOTO_MAX + 1)
    if not data:
        raise HTTPException(422, "فایل خالی است")
    if len(data) > PHOTO_MAX:
        raise HTTPException(413, "حجم عکس بیشتر از ۱٫۵ مگابایت است")
    mime = next((m for sig, m in IMAGE_MAGIC if data.startswith(sig) and (m != "image/webp" or data[8:12] == b"WEBP")), None)
    if mime is None:
        raise HTTPException(422, "فقط عکس JPG، PNG یا WEBP پذیرفته می‌شود")
    key = _photo_key(pid)
    existing = db.scalars(select(BotFile).where(BotFile.bot_id == bot_id, BotFile.block_id == key)).first()
    photos = db.scalar(select(func.coalesce(func.sum(BotFile.size), 0)).where(BotFile.bot_id == bot_id, BotFile.block_id.like("product:%"))) or 0
    if photos - (existing.size if existing else 0) + len(data) > 4 * MAX_TOTAL:  # their own allowance, apart from message files
        raise HTTPException(413, "فضای عکس‌های این ربات پر شده است؛ عکس چند محصول را حذف کنید")
    ext = {"image/jpeg": "jpg", "image/png": "png", "image/webp": "webp"}[mime]
    if existing is None:
        db.add(BotFile(bot_id=bot_id, block_id=key, kind="image", filename=f"product-{pid}.{ext}", mime=mime, size=len(data), data=data))
    else:
        existing.filename, existing.mime, existing.size, existing.data = f"product-{pid}.{ext}", mime, len(data), data
    p.has_photo = True
    db.commit()
    return _prod_out(p)


@router.get("/api/bots/{bot_id}/catalog/products/{pid}/photo")
def get_photo(bot_id: int, pid: int, user: User = Depends(current_user), db: Session = Depends(get_db)):
    from fastapi import Response
    _own(bot_id, user, db)
    _product(bot_id, pid, db)
    f = db.scalars(select(BotFile).where(BotFile.bot_id == bot_id, BotFile.block_id == _photo_key(pid))).first()
    if f is None:
        raise HTTPException(404, "این محصول عکس ندارد")
    return Response(f.data, media_type=f.mime, headers={"Cache-Control": "private, max-age=60"})


@router.delete("/api/bots/{bot_id}/catalog/products/{pid}/photo")
def delete_photo(bot_id: int, pid: int, user: User = Depends(current_user), db: Session = Depends(get_db)):
    _own(bot_id, user, db)
    p = _product(bot_id, pid, db)
    _drop_photos(db, bot_id, [pid])
    p.has_photo = False
    db.commit()
    return _prod_out(p)
